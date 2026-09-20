from openpyxl import load_workbook

from app.exporters.excel.service import exportar_carrera
from app.models.carrera import Carrera
from app.models.enums import ModalidadPrograma, TipoElemento, TipoReconocimiento
from app.models.plan_curricular import PlanCurricular
from app.models.semestre import Semestre
from app.models.semestre_elemento import SemestreElemento
from app.models.materia import Materia
from app.models.enums import TipoMateria
from app.services import plan_curricular as plan_service
from app.services import configuracion_reglas as config_service
from app.schemas.plan_curricular import PlanCurricularCreate, PlanCurricularUpdate


def test_recalculo_de_claves_es_explicito_y_salta_espacios_optativos(db_session):
    carrera = Carrera(clave="KEY", nombre="Claves")
    db_session.add(carrera)
    db_session.flush()
    plan = PlanCurricular(carrera_id=carrera.id, clave="KEY-2026", mnemonico="KEY", decanato_otro="Otro", max_creditos_semestre=50)
    db_session.add(plan)
    db_session.flush()
    s1 = Semestre(plan_curricular_id=plan.id, numero=1)
    s2 = Semestre(plan_curricular_id=plan.id, numero=2)
    db_session.add_all([s1, s2])
    db_session.flush()
    materias = [Materia(plan_curricular_id=plan.id, clave=f"OLD{i}", nombre=f"Materia {i}", horas_docente=32, horas_independientes=64, tipo=TipoMateria.OBLIGATORIA) for i in range(1, 5)]
    db_session.add_all(materias)
    db_session.flush()
    db_session.add_all([
        SemestreElemento(semestre_id=s1.id, tipo=TipoElemento.MATERIA, materia_id=materias[0].id, orden=0),
        SemestreElemento(semestre_id=s1.id, tipo=TipoElemento.ESPACIO_OPTATIVO, nombre="Optativa", horas_docente=32, horas_independientes=64, orden=1),
        SemestreElemento(semestre_id=s1.id, tipo=TipoElemento.MATERIA, materia_id=materias[3].id, orden=2),
        SemestreElemento(semestre_id=s2.id, tipo=TipoElemento.MATERIA, materia_id=materias[1].id, orden=0),
        SemestreElemento(semestre_id=s2.id, tipo=TipoElemento.MATERIA, materia_id=materias[2].id, orden=1),
    ])
    db_session.commit()

    plan_service.recalcular_claves(db_session, plan.id)
    assert [m.clave for m in materias] == ["KEY001", "KEY002", "KEY003", "KEY004"]


def test_exportacion_condicional_flexible_y_optativas(db_session):
    carrera = Carrera(clave="EXP", nombre="Exportación")
    db_session.add(carrera)
    db_session.flush()
    flexible = PlanCurricular(carrera_id=carrera.id, clave="EXP-F", mnemonico="EXF", decanato_otro="Otro", modalidad=ModalidadPrograma.NO_ESCOLARIZADA, tipo_reconocimiento=TipoReconocimiento.FEDERAL, max_creditos_semestre=50)
    estatal = PlanCurricular(carrera_id=carrera.id, clave="EXP-E", mnemonico="EXE", decanato_otro="Otro", modalidad=ModalidadPrograma.MIXTA, tipo_reconocimiento=TipoReconocimiento.ESTATAL, max_creditos_semestre=50)
    db_session.add_all([flexible, estatal])
    db_session.commit()

    wb_flexible = load_workbook(exportar_carrera(db_session, flexible.id))
    assert "Flexible" in wb_flexible.sheetnames
    assert "Optativas" not in wb_flexible.sheetnames

    wb_estatal = load_workbook(exportar_carrera(db_session, estatal.id))
    assert "Optativas" in wb_estatal.sheetnames
    assert "Flexible" not in wb_estatal.sheetnames
    hoja = wb_estatal["Optativas"]
    assert hoja.cell(row=2, column=2).value == "Inteligencia emocional"
    assert hoja.cell(row=19, column=2).value == "Visualización de datos empresariales"


def test_cambio_de_programa_genera_nueva_version_de_reglas(db_session):
    carrera = Carrera(clave="CFG", nombre="Configuración")
    db_session.add(carrera)
    db_session.flush()
    plan = plan_service.create_plan(
        db_session,
        carrera.id,
        PlanCurricularCreate(clave="CFG-2026", mnemonico="CFG"),
    )
    plan_service.update_plan(
        db_session,
        plan.id,
        PlanCurricularUpdate(modalidad=ModalidadPrograma.NO_ESCOLARIZADA),
    )
    versiones = config_service.listar(db_session, plan.id)
    assert [v.version for v in versiones] == [1, 2]
    assert versiones[0].vigente is False
    assert versiones[1].vigente is True
    assert versiones[1].max_materias_ciclo == 5
    assert versiones[1].prohibir_seriacion is True
