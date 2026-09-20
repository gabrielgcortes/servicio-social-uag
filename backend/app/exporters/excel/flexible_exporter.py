"""Hoja Flexible agrupada por área, sin seriaciones."""
from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

from app.models.enums import TipoElemento

_COLUMNAS = ["Área de formación", "Clave", "Nombre", "Horas docente", "Horas independientes", "Créditos", "Instalación", "Modalidad", "Docente sugerido", "Aporte sustancial", "Programa de asignatura"]


def escribir_flexible(hoja: Worksheet, mapa: dict) -> None:
    for col, titulo in enumerate(_COLUMNAS, 1):
        hoja.cell(row=1, column=col, value=titulo)
    materias = []
    for semestre in mapa["semestres"]:
        for elemento in semestre["elementos"]:
            if elemento.tipo == TipoElemento.MATERIA and elemento.materia is not None:
                materias.append(elemento.materia)
    materias.sort(key=lambda m: ((m.area_formacion.value if m.area_formacion else "SIN_AREA"), m.nombre))
    fila = 2
    for materia in materias:
        valores = [
            materia.area_formacion.value if materia.area_formacion else "Sin área",
            materia.clave, materia.nombre, materia.horas_docente, materia.horas_independientes,
            float(materia.creditos or 0), materia.tipo_aula.value if materia.tipo_aula else (materia.instalaciones or ""),
            materia.modalidad.value if materia.modalidad else "", materia.docente_sugerido or "",
            "Sí" if materia.aporte_sustancial else "No", materia.programa_asignatura or "",
        ]
        for col, valor in enumerate(valores, 1):
            hoja.cell(row=fila, column=col, value=valor)
        fila += 1

    total_hd = sum(s["totales"]["total_horas_docente"] for s in mapa["semestres"])
    total_hi = sum(s["totales"]["total_horas_independientes"] for s in mapa["semestres"])
    total_creditos = sum(s["totales"]["total_creditos"] for s in mapa["semestres"])
    obligatorias = len(materias)
    plan = mapa["plan"]
    base = plan.texto_administrativo_flexible or (
        "El plan de estudios se diseñó bajo una trayectoria curricular flexible. "
        "La organización de los contenidos es congruente con el perfil de egreso y no se tienen seriaciones.\n\n"
        "El plan se integra por {asignaturas} asignaturas obligatorias, {horas} horas de aprendizaje "
        "({horas_docente} con académico y {horas_independientes} independientes), equivalentes a "
        "{creditos:g} créditos."
    )
    texto = base.format(asignaturas=obligatorias, horas=total_hd + total_hi, horas_docente=total_hd, horas_independientes=total_hi, creditos=total_creditos)
    fila += 2
    hoja.cell(row=fila, column=1, value="Administración del plan de estudios")
    hoja.cell(row=fila + 1, column=1, value=texto)
    hoja.merge_cells(start_row=fila + 1, start_column=1, end_row=fila + 4, end_column=len(_COLUMNAS))
    hoja.cell(row=fila + 1, column=1).alignment = __import__("openpyxl").styles.Alignment(wrap_text=True, vertical="top")
