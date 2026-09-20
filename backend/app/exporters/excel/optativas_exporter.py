"""Escribe la hoja "Optativas" a partir del catálogo de materias optativas de
la carrera."""
from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

from app.exporters.excel.layout import COLUMNAS_OPTATIVAS, FILA_INICIAL_OPTATIVAS
from app.models.materia import Materia

def escribir_optativas(hoja: Worksheet, optativas: list[Materia], *, institucionales: list[str] | None = None, ciclos_institucionales: list[int] | None = None) -> None:
    fila = FILA_INICIAL_OPTATIVAS
    for columna, encabezado in enumerate(COLUMNAS_OPTATIVAS, start=1):
        hoja.cell(row=fila, column=columna, value=encabezado)
    fila += 1

    por_id = {m.id: m for m in optativas}
    for materia in optativas:
        prerequisito = por_id.get(materia.seriacion_materia_id)
        valores = [
            materia.clave,
            materia.nombre,
            materia.horas_docente,
            materia.horas_independientes,
            float(materia.creditos) if materia.creditos is not None else "",
            materia.instalaciones or "",
            materia.modalidad.value if materia.modalidad else "",
            prerequisito.clave if prerequisito else "",
            materia.area_formacion.value if materia.area_formacion else "",
            materia.docente_sugerido or "",
            "Sí" if materia.aporte_sustancial else "No",
            materia.programa_asignatura or "",
            ", ".join(str(c) for c in materia.ciclos_disponibles),
        ]
        for columna, valor in enumerate(valores, start=1):
            hoja.cell(row=fila, column=columna, value=valor)
        fila += 1
    if institucionales:
        ciclos = ", ".join(str(c) for c in sorted(set(ciclos_institucionales or [])))
        for nombre in institucionales:
            valores = ["", nombre, "", "", "", "", "", "", "", "", "", "", ciclos]
            for columna, valor in enumerate(valores, start=1):
                hoja.cell(row=fila, column=columna, value=valor)
            fila += 1
