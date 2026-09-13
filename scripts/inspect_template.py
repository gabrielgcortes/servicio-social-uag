"""Inspecciona templates/mapa_curricular.xlsx para ayudar a construir
exporters/excel/layout.py con las coordenadas reales de la plantilla
institucional.

Este script NO se ejecuta automáticamente como parte de la implementación:
el usuario debe correrlo manualmente (requiere `pip install openpyxl`) y
compartir la salida para afinar el layout de la exportación.

Uso:
    python scripts/inspect_template.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "templates" / "mapa_curricular.xlsx"


def main() -> None:
    if not TEMPLATE_PATH.exists():
        print(f"No se encontró la plantilla en {TEMPLATE_PATH}")
        sys.exit(1)

    wb = load_workbook(TEMPLATE_PATH, data_only=False)
    print(f"Hojas: {wb.sheetnames}\n")

    for nombre_hoja in wb.sheetnames:
        hoja = wb[nombre_hoja]
        print(f"=== Hoja: {nombre_hoja} ===")
        print(f"Dimensiones: {hoja.dimensions}")
        print(f"Rangos combinados ({len(hoja.merged_cells.ranges)}):")
        for rango in hoja.merged_cells.ranges:
            print(f"  {rango}")

        print("Celdas con valor en las primeras 25 filas x 15 columnas:")
        max_fila = min(hoja.max_row, 25)
        max_columna = min(hoja.max_column, 15)
        for fila in range(1, max_fila + 1):
            for columna in range(1, max_columna + 1):
                celda = hoja.cell(row=fila, column=columna)
                if celda.value is not None:
                    print(f"  [{celda.coordinate}] = {celda.value!r}")
        print()


if __name__ == "__main__":
    main()
